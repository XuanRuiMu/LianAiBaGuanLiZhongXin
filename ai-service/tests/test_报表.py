import io
import asyncio
import os

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import app.api.报表 as 报表路由
import app.报表.定时 as 定时模块
from app.报表.生成器 import 生成Excel, 生成PDF

样例概览 = {"totalUsers": 120, "totalRoles": 200, "totalMessages": 8000}
样例趋势 = [
    {"date": "2026-09-05", "newUsers": 5, "totalUsers": 120},
    {"date": "2026-09-06", "newUsers": 3, "totalUsers": 123},
]


async def 假数据():
    return 样例概览, 样例趋势


class Test生成器:
    def test_excel可打开且含数据(self, tmp_path):
        数据 = 生成Excel(样例概览, 样例趋势)
        assert 数据[:2] == b"PK"
        路径 = tmp_path / "日报.xlsx"
        路径.write_bytes(数据)
        簿 = load_workbook(路径)
        页 = 簿.active
        文本 = " ".join(str(格.value) for 行 in 页.iter_rows() for 格 in 行)
        assert "totalUsers" in 文本 and "2026-09-05" in 文本

    def test_pdf文件头正确(self):
        数据 = 生成PDF(样例概览, 样例趋势)
        assert 数据[:5] == b"%PDF-"
        assert len(数据) > 1000


class Test定时落盘:
    @pytest.mark.asyncio
    async def test_生成并落盘双格式(self, monkeypatch, tmp_path):
        import app.config as 配置模块
        from app.config import 取配置
        monkeypatch.setattr(定时模块, "取日报数据", 假数据)
        monkeypatch.setenv("REPORT_DIR", str(tmp_path))
        monkeypatch.setenv("REPORT_INTERVAL_SECONDS", "3600")
        取配置.cache_clear()
        try:
            路径表 = await 定时模块.生成并落盘()
            assert set(路径表) == {"xlsx", "pdf"}
            assert load_workbook(路径表["xlsx"]).active.max_row > 5
            assert open(路径表["pdf"], "rb").read(5) == b"%PDF-"
        finally:
            取配置.cache_clear()

    @pytest.mark.asyncio
    async def test_定时循环遇错不崩且可停止(self, monkeypatch):
        async def 必失败():
            raise RuntimeError("后端未就绪")
        monkeypatch.setattr(定时模块, "生成并落盘", 必失败)
        停止事件 = asyncio.Event()
        任务 = asyncio.create_task(定时模块.报表定时循环(停止事件))
        await asyncio.sleep(0.1)
        停止事件.set()
        await asyncio.wait_for(任务, timeout=5)
        assert 任务.done()


class Test下载端点:
    def test_无凭证拒收(self):
        from app.main import app
        客户端 = TestClient(app, raise_server_exceptions=False)
        assert 客户端.get("/api/v1/reports/日报.xlsx").status_code == 401
        assert 客户端.get("/api/v1/reports/日报.pdf").status_code == 401

    def test_凭证通过可下载(self, monkeypatch, tmp_path):
        import app.config as 配置模块
        from app.config import 取配置
        monkeypatch.setattr(报表路由, "取日报数据", 假数据)
        monkeypatch.setenv("REPORT_DIR", str(tmp_path))
        取配置.cache_clear()
        try:
            from app.main import app
            客户端 = TestClient(app, raise_server_exceptions=False)
            头 = {"X-Internal-Token": os.environ.get("INTERNAL_TOKEN", "test-token")}
            响应 = 客户端.get("/api/v1/reports/日报.xlsx", headers=头)
            assert 响应.status_code == 200
            assert 响应.headers["content-type"].startswith(
                "application/vnd.openxmlformats")
            簿 = load_workbook(io.BytesIO(响应.content))
            assert 簿.active.max_row > 5
            响应二 = 客户端.get("/api/v1/reports/日报.pdf", headers=头)
            assert 响应二.status_code == 200
            assert 响应二.content[:5] == b"%PDF-"
        finally:
            取配置.cache_clear()
